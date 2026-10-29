#!/usr/bin/env python3
"""
美纳里尼合规报销智能助手Web应用版本 v2.0
优化版 - 更美观的界面和更好的排版
"""

from flask import Flask, render_template, request, jsonify, send_from_directory
import os
import json
import re
from typing import List, Dict, Any
import threading
import time
from datetime import datetime
import werkzeug.utils

app = Flask(__name__)

# 初始化知识库数据
qa_pairs = []
initialized = False

# 配置参数
MAX_ANSWERS = 2
SIMILARITY_THRESHOLD = 0.4

class SearchEngine:
    """搜索引擎"""

    def __init__(self):
        pass

    def parse_qa_pairs(self) -> List[Dict[str, Any]]:
        """解析QA对"""
        try:
            from docx import Document
            doc = Document("C:\\\\Users\\\\jj\\\\Desktop\\\\RAG\\\\QA.docx")
            qa_pairs = []

            for i in range(len(doc.paragraphs) - 1):
                para1 = doc.paragraphs[i]
                para2 = doc.paragraphs[i + 1]

                text1 = para1.text.strip()
                text2 = para2.text.strip()

                if text1 and ('？' in text1 or '?' in text1) and len(text1) > 5:
                    if text2 and len(text2) > 5:
                        qa = {
                            'id': len(qa_pairs) + 1,
                            'question': text1,
                            'answer': text2,
                            'index': i + 1
                        }
                        qa_pairs.append(qa)

            return qa_pairs
        except Exception as e:
            print(f"解析失败: {e}")
            return []

    def search_with_fuzzy(self, query: str, qa_pairs: List[Dict[str, Any]], max_answers: int = MAX_ANSWERS) -> List[Dict[str, Any]]:
        """模糊搜索"""
        query_lower = query.lower()
        results = []

        # 定义同义词词典
        synonyms = {
            '重置密码': ['重置密码', '重置口令', '找回密码', '修改密码'],
            '邮箱': ['邮箱', '电子邮箱', 'email', '邮件'],
            '收不到': ['收不到', '未收到', '没有收到', '没收到'],
            '登录': ['登录', '登入', 'signin', 'log in'],
            '报销': ['报销', '报销流程', '报销申请', '报销系统']
        }

        # 扩展查询词
        expanded_terms = set()
        for term, synonyms_list in synonyms.items():
            if term in query_lower:
                expanded_terms.update(synonyms_list)

        # 如果没有找到同义词，使用原查询
        if not expanded_terms:
            expanded_terms.add(query_lower)

        for qa in qa_pairs:
            score = 0
            q_text = qa['question'].lower()
            a_text = qa['answer'].lower()

            # 计算每个扩展词的匹配度
            for term in expanded_terms:
                if term in q_text:
                    score += 1.0
                elif term in a_text:
                    score += 0.7
                else:
                    # 部分匹配
                    words = term.split()
                    match_count = sum(1 for word in words if word in q_text or word in a_text)
                    if match_count > 0:
                        score += match_count / len(words) * 0.5

            if score > 0:
                results.append({
                    'id': qa['id'],
                    'question': qa['question'],
                    'answer': qa['answer'],
                    'similarity': score / len(expanded_terms),
                    'index': qa.get('index', 0)
                })

        # 按相似度排序
        results.sort(key=lambda x: x['similarity'], reverse=True)
        return results[:max_answers]

    def highlight_keywords(self, text: str, keywords: List[str]) -> str:
        """高亮关键词"""
        if not keywords:
            return text

        for keyword in keywords:
            if keyword:
                text = re.sub(
                    re.escape(keyword),
                    f'<mark class="highlight-mark">{keyword}</mark>',
                    text,
                    flags=re.IGNORECASE
                )
        return text

    def extract_keywords(self, text: str) -> List[str]:
        """提取关键词"""
        words = re.findall(r'[\\u4e00-\\u9fff]+|[a-zA-Z]+', text)
        stopwords = {'的', '了', '在', '是', '我', '有', '和', '就', '不', '人', '都', '一', '个', '上', '也', '很', '到', '说', '要', '去', '你', '会', '着', '没有', '看', '好', '自己', '这'}
        keywords = []
        for word in words:
            if len(word) >= 2 and word not in stopwords and word not in keywords:
                keywords.append(word)
        return keywords[:5]

# 全局搜索引擎实例
search_engine = SearchEngine()

def load_knowledge_base():
    """加载知识库"""
    global qa_pairs, initialized
    if not initialized:
        try:
            qa_pairs = search_engine.parse_qa_pairs()
            initialized = True
            print("知识库加载成功，共 %d 个问答对" % len(qa_pairs))
        except Exception as e:
            print(f"知识库加载失败: {e}")

# 路由定义
@app.route('/')
def index():
    """主页"""
    # 如果HTML文件存在，使用它
    if os.path.exists('web_app_v2.html'):
        return send_from_directory('.', 'web_app_v2.html')
    return render_template('index_v2.html')

@app.route('/api/config', methods=['GET'])
def get_config():
    """获取配置"""
    return jsonify({
        'max_answers': MAX_ANSWERS,
        'similarity_threshold': SIMILARITY_THRESHOLD
    })

@app.route('/api/config', methods=['POST'])
def update_config():
    """更新配置"""
    try:
        data = request.json
        global MAX_ANSWERS, SIMILARITY_THRESHOLD
        if 'max_answers' in data:
            MAX_ANSWERS = int(data['max_answers'])
        if 'similarity_threshold' in data:
            SIMILARITY_THRESHOLD = float(data['similarity_threshold'])
        return jsonify({'success': True})
    except Exception as e:
        return jsonify({'error': str(e)}), 400

@app.route('/api/upload', methods=['POST'])
def upload_knowledge_base():
    """上传知识库"""
    try:
        if 'file' not in request.files:
            return jsonify({'error': '没有文件上传'}), 400

        file = request.files['file']
        if file.filename == '':
            return jsonify({'error': '没有选择文件'}), 400

        if file and file.filename.endswith('.docx'):
            # 保存上传的文件
            upload_path = os.path.join('uploads', file.filename)
            os.makedirs('uploads', exist_ok=True)
            file.save(upload_path)

            # 重新解析知识库
            global qa_pairs, initialized
            qa_pairs = []
            initialized = False

            # 解析新上传的文件
            try:
                from docx import Document
                doc = Document(upload_path)
                qa_pairs = []

                for i in range(len(doc.paragraphs) - 1):
                    para1 = doc.paragraphs[i]
                    para2 = doc.paragraphs[i + 1]

                    text1 = para1.text.strip()
                    text2 = para2.text.strip()

                    if text1 and ('？' in text1 or '?' in text1) and len(text1) > 5:
                        if text2 and len(text2) > 5:
                            qa = {
                                'id': len(qa_pairs) + 1,
                                'question': text1,
                                'answer': text2,
                                'index': i + 1
                            }
                            qa_pairs.append(qa)

                initialized = True
                return jsonify({
                    'success': True,
                    'message': '知识库更新成功',
                    'qa_pairs_count': len(qa_pairs)
                })
            except Exception as e:
                return jsonify({'error': f'解析失败: {str(e)}'}), 500
        else:
            return jsonify({'error': '只支持.docx文件'}), 400
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/ask', methods=['POST'])
def ask():
    """问答API"""
    try:
        data = request.json
        question = data.get('question', '').strip()

        # 获取配置参数
        max_answers = data.get('max_answers', MAX_ANSWERS)
        similarity_threshold = data.get('similarity_threshold', SIMILARITY_THRESHOLD)

        if not question:
            return jsonify({'error': '请输入问题'}), 400

        # 加载知识库
        if not initialized:
            load_knowledge_base()

        if not initialized:
            return jsonify({'error': '知识库未加载'}), 500

        # 搜索答案
        results = search_engine.search_with_fuzzy(question, qa_pairs, max_answers=max_answers)

        if results:
            # 提取关键词并高亮
            keywords = search_engine.extract_keywords(question)

            # 按相似度分组
            high_similarity = [r for r in results if r['similarity'] >= SIMILARITY_THRESHOLD]
            low_similarity = [r for r in results if r['similarity'] < SIMILARITY_THRESHOLD and r['similarity'] > 0]

            # 构建回答
            response_data = {
                'success': True,
                'high_similarity': [],
                'low_similarity': [],
                'keywords': keywords
            }

            if high_similarity:
                for i, result in enumerate(high_similarity, 1):
                    answer = search_engine.highlight_keywords(result['answer'], keywords)
                    response_data['high_similarity'].append({
                        'id': i,
                        'answer': answer,
                        'similarity': result['similarity'],
                        'original_question': result['question']
                    })

            if low_similarity and len(results) > len(high_similarity):
                for i, result in enumerate(low_similarity[:3], 1):
                    answer = search_engine.highlight_keywords(result['answer'], keywords)
                    response_data['low_similarity'].append({
                        'id': i + len(high_similarity),
                        'answer': answer,
                        'similarity': result['similarity'],
                        'original_question': result['question']
                    })

            return jsonify(response_data)
        else:
            return jsonify({
                'success': False,
                'error': '未找到相关答案',
                'message': '''很抱歉，我在知识库中没有找到与您问题相关的信息。

📝 建议：
1. 请检查问题表述是否清晰
2. 尝试使用不同的关键词描述您的问题
3. 如果仍需帮助，请联系人工客服

📞 客服联系方式：请通过公司内部通讯系统联系IT支持团队'''
            })

    except Exception as e:
        return jsonify({'error': f'服务器错误: {str(e)}'}), 500

@app.route('/api/status')
def status():
    """获取状态"""
    return jsonify({
        'status': 'ok',
        'initialized': initialized,
        'qa_pairs_count': len(qa_pairs),
        'timestamp': datetime.now().isoformat()
    })

@app.route('/static/<path:filename>')
def static_files(filename):
    """静态文件"""
    return send_from_directory('static', filename)

# 创建模板和静态文件
def create_files():
    """创建模板和静态文件"""
    os.makedirs('templates', exist_ok=True)
    os.makedirs('static', exist_ok=True)
    os.makedirs('static/css', exist_ok=True)
    os.makedirs('static/js', exist_ok=True)

    # 创建CSS文件
    css_content = '''
/* 全局样式 */
:root {
    --primary-color: #667eea;
    --secondary-color: #764ba2;
    --success-color: #28a745;
    --warning-color: #ffc107;
    --danger-color: #dc3545;
    --light-color: #f8f9fa;
    --dark-color: #343a40;
    --border-radius: 15px;
    --transition: all 0.3s ease;
}

* {
    margin: 0;
    padding: 0;
    box-sizing: border-box;
}

body {
    font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, 'Helvetica Neue', Arial, sans-serif;
    background: linear-gradient(135deg, var(--primary-color) 0%, var(--secondary-color) 100%);
    min-height: 100vh;
    display: flex;
    align-items: center;
    justify-content: center;
    padding: 20px;
    line-height: 1.6;
}

/* 主容器 */
.chat-container {
    max-width: 900px;
    width: 100%;
    height: 100vh;
    display: flex;
    flex-direction: column;
    gap: 20px;
}

/* 头部 */
.header {
    background: rgba(255, 255, 255, 0.95);
    backdrop-filter: blur(10px);
    border-radius: var(--border-radius);
    padding: 25px;
    box-shadow: 0 10px 30px rgba(0,0,0,0.1);
    display: flex;
    align-items: center;
    gap: 20px;
    animation: slideDown 0.5s ease-out;
}

@keyframes slideDown {
    from {
        opacity: 0;
        transform: translateY(-20px);
    }
    to {
        opacity: 1;
        transform: translateY(0);
    }
}

.header-icon {
    font-size: 3rem;
    animation: float 3s ease-in-out infinite;
}

@keyframes float {
    0%, 100% { transform: translateY(0); }
    50% { transform: translateY(-10px); }
}

.header-content h1 {
    font-size: 2.2rem;
    color: var(--primary-color);
    margin-bottom: 5px;
    font-weight: 700;
}

.header-content p {
    color: #6c757d;
    font-size: 1rem;
}

/* 消息区域 */
.messages-container {
    flex: 1;
    overflow: hidden;
    display: flex;
    flex-direction: column;
}

.messages {
    flex: 1;
    overflow-y: auto;
    padding: 25px;
    background: rgba(255, 255, 255, 0.9);
    backdrop-filter: blur(10px);
    border-radius: var(--border-radius);
    box-shadow: 0 10px 30px rgba(0,0,0,0.1);
    display: flex;
    flex-direction: column;
    gap: 15px;
}

.messages::-webkit-scrollbar {
    width: 8px;
}

.messages::-webkit-scrollbar-track {
    background: #f1f1f1;
    border-radius: 10px;
}

.messages::-webkit-scrollbar-thumb {
    background: var(--primary-color);
    border-radius: 10px;
}

/* 消息样式 */
.message {
    animation: fadeInUp 0.3s ease-out;
    position: relative;
    padding: 20px;
    border-radius: var(--border-radius);
    max-width: 85%;
    word-wrap: break-word;
}

@keyframes fadeInUp {
    from {
        opacity: 0;
        transform: translateY(20px);
    }
    to {
        opacity: 1;
        transform: translateY(0);
    }
}

.message.user {
    background: linear-gradient(135deg, var(--primary-color), var(--secondary-color));
    color: white;
    margin-left: auto;
    align-self: flex-end;
    box-shadow: 0 5px 15px rgba(102, 126, 234, 0.3);
}

.message.assistant {
    background: white;
    color: #333;
    border-left: 4px solid var(--primary-color);
    box-shadow: 0 3px 10px rgba(0,0,0,0.1);
}

.message-header {
    display: flex;
    justify-content: space-between;
    align-items: center;
    margin-bottom: 10px;
    font-size: 0.85rem;
    opacity: 0.7;
}

.message-content {
    line-height: 1.8;
}

/* 回答卡片 */
.answer-card {
    background: white;
    border-radius: 10px;
    padding: 20px;
    margin: 10px 0;
    box-shadow: 0 2px 10px rgba(0,0,0,0.05);
    border: 1px solid #e9ecef;
    transition: var(--transition);
}

.answer-card:hover {
    box-shadow: 0 5px 20px rgba(0,0,0,0.1);
    transform: translateY(-2px);
}

.answer-card.high {
    border-left: 4px solid #28a745;
}

.answer-card.medium {
    border-left: 4px solid #ffc107;
}

.answer-title {
    display: flex;
    align-items: center;
    gap: 10px;
    margin-bottom: 15px;
    font-weight: 600;
}

.answer-title-icon {
    width: 30px;
    height: 30px;
    border-radius: 50%;
    display: flex;
    align-items: center;
    justify-content: center;
    font-size: 14px;
}

.answer-title.high .answer-title-icon {
    background: #d4edda;
    color: #155724;
}

.answer-title.medium .answer-title-icon {
    background: #fff3cd;
    color: #856404;
}

.similarity-badge {
    background: var(--primary-color);
    color: white;
    padding: 3px 10px;
    border-radius: 15px;
    font-size: 0.8rem;
    font-weight: 500;
}

.similarity-badge.high {
    background: #28a745;
}

.similarity-badge.medium {
    background: #ffc107;
    color: #212529;
}

/* 输入区域 */
.input-container {
    background: rgba(255, 255, 255, 0.95);
    backdrop-filter: blur(10px);
    border-radius: var(--border-radius);
    padding: 25px;
    box-shadow: 0 10px 30px rgba(0,0,0,0.1);
}

.input-group {
    display: flex;
    gap: 15px;
    align-items: center;
}

.form-control {
    flex: 1;
    border: 2px solid #e9ecef;
    border-radius: 25px;
    padding: 12px 25px;
    font-size: 1rem;
    transition: var(--transition);
}

.form-control:focus {
    border-color: var(--primary-color);
    box-shadow: 0 0 0 0.2rem rgba(102, 126, 234, 0.25);
    outline: none;
}

.btn-primary {
    border-radius: 25px;
    padding: 12px 30px;
    background: linear-gradient(135deg, var(--primary-color), var(--secondary-color));
    border: none;
    color: white;
    font-weight: 600;
    transition: var(--transition);
    white-space: nowrap;
}

.btn-primary:hover {
    transform: translateY(-2px);
    box-shadow: 0 5px 20px rgba(102, 126, 234, 0.4);
}

/* 输入状态 */
.typing-indicator {
    display: inline-flex;
    align-items: center;
    gap: 4px;
    padding: 10px;
}

.typing-indicator span {
    height: 8px;
    width: 8px;
    background-color: var(--primary-color);
    border-radius: 50%;
    display: inline-block;
    animation: bounce 1.4s infinite ease-in-out both;
}

.typing-indicator span:nth-child(1) {
    animation-delay: -0.32s;
}

.typing-indicator span:nth-child(2) {
    animation-delay: -0.16s;
}

@keyframes bounce {
    0%, 80%, 100% {
        transform: scale(0);
    }
    40% {
        transform: scale(1);
    }
}

/* 关键词高亮 */
.highlight-mark {
    background: linear-gradient(120deg, #f6d365 0%, #fda085 100%);
    color: white;
    padding: 2px 6px;
    border-radius: 4px;
    font-weight: 500;
    box-shadow: 0 2px 4px rgba(0,0,0,0.1);
}

/* 状态指示器 */
.status-badge {
    position: fixed;
    top: 20px;
    right: 20px;
    padding: 10px 20px;
    background: rgba(255, 255, 255, 0.95);
    backdrop-filter: blur(10px);
    border-radius: 20px;
    font-size: 14px;
    box-shadow: 0 5px 15px rgba(0,0,0,0.1);
    display: flex;
    align-items: center;
    gap: 8px;
    z-index: 1000;
}

.status-badge.online {
    color: var(--success-color);
}

.status-badge.offline {
    color: var(--danger-color);
}

.status-badge .status-dot {
    width: 8px;
    height: 8px;
    border-radius: 50%;
    background: currentColor;
    animation: pulse 2s infinite;
}

@keyframes pulse {
    0% {
        box-shadow: 0 0 0 0 currentColor;
    }
    70% {
        box-shadow: 0 0 0 10px rgba(255,255,255,0);
    }
    100% {
        box-shadow: 0 0 0 0 rgba(255,255,255,0);
    }
}

/* 快捷问题 */
.quick-questions {
    display: flex;
    flex-wrap: wrap;
    gap: 10px;
    margin-top: 15px;
}

.quick-question {
    background: #f8f9fa;
    border: 1px solid #e9ecef;
    border-radius: 20px;
    padding: 8px 16px;
    font-size: 0.9rem;
    color: #495057;
    cursor: pointer;
    transition: var(--transition);
}

.quick-question:hover {
    background: var(--primary-color);
    color: white;
    border-color: var(--primary-color);
    transform: translateY(-2px);
}

/* 响应式设计 */
@media (max-width: 768px) {
    .chat-container {
        height: 100vh;
    }

    .header {
        flex-direction: column;
        text-align: center;
        padding: 20px;
    }

    .header-icon {
        font-size: 2.5rem;
    }

    .header-content h1 {
        font-size: 1.8rem;
    }

    .message {
        max-width: 90%;
    }

    .input-group {
        flex-direction: column;
    }

    .form-control {
        width: 100%;
    }

    .btn-primary {
        width: 100%;
    }

    .status-badge {
        top: 10px;
        right: 10px;
        font-size: 12px;
        padding: 8px 16px;
    }
}

/* 加载动画 */
.loading-spinner {
    display: inline-block;
    width: 20px;
    height: 20px;
    border: 3px solid #f3f3f3;
    border-top: 3px solid var(--primary-color);
    border-radius: 50%;
    animation: spin 1s linear infinite;
}

@keyframes spin {
    0% { transform: rotate(0deg); }
    100% { transform: rotate(360deg); }
}
'''

    with open('static/css/style.css', 'w', encoding='utf-8') as f:
        f.write(css_content)

    # 创建JS文件
    js_content = '''
// 聊天应用主逻辑
class ChatApp {
    constructor() {
        this.messagesContainer = document.getElementById('messages');
        this.chatForm = document.getElementById('chatForm');
        this.userInput = document.getElementById('userInput');
        this.sendBtn = document.getElementById('sendBtn');
        this.statusBadge = document.getElementById('statusBadge');
        this.isTyping = false;

        this.init();
    }

    init() {
        // 检查连接状态
        this.checkStatus();
        setInterval(() => this.checkStatus(), 5000);

        // 绑定事件
        this.chatForm.addEventListener('submit', (e) => {
            e.preventDefault();
            this.sendMessage();
        });

        this.userInput.addEventListener('keypress', (e) => {
            if (e.key === 'Enter' && !e.shiftKey) {
                e.preventDefault();
                this.sendMessage();
            }
        });

        // 快捷问题点击
        document.querySelectorAll('.quick-question').forEach(btn => {
            btn.addEventListener('click', () => {
                const question = btn.textContent;
                this.userInput.value = question;
                this.sendMessage();
            });
        });

        // 页面加载完成后聚焦输入框
        document.addEventListener('DOMContentLoaded', () => {
            this.userInput.focus();
        });
    }

    async checkStatus() {
        try {
            const response = await fetch('/api/status');
            const data = await response.json();

            if (data.status === 'ok') {
                this.statusBadge.innerHTML = `
                    <span class="status-dot"></span>
                    <span>在线 (${data.qa_pairs_count}个问答)</span>
                `;
                this.statusBadge.classList.add('online');
                this.statusBadge.classList.remove('offline');
            }
        } catch (error) {
            this.statusBadge.innerHTML = `
                <span class="status-dot"></span>
                <span>离线</span>
            `;
            this.statusBadge.classList.add('offline');
            this.statusBadge.classList.remove('online');
        }
    }

    addMessage(content, type) {
        const messageDiv = document.createElement('div');
        messageDiv.className = `message ${type}`;

        if (type === 'user') {
            messageDiv.innerHTML = `
                <div class="message-content">${content}</div>
            `;
        } else {
            messageDiv.innerHTML = `
                <div class="message-header">
                    <span>🤖 合规助手</span>
                    <span>${new Date().toLocaleTimeString()}</span>
                </div>
                <div class="message-content"></div>
            `;
            messageDiv.querySelector('.message-content').innerHTML = content;
        }

        this.messagesContainer.appendChild(messageDiv);
        this.scrollToBottom();

        // 添加动画效果
        messageDiv.style.opacity = '0';
        setTimeout(() => {
            messageDiv.style.opacity = '1';
        }, 100);
    }

    showTyping() {
        if (this.isTyping) return;
        this.isTyping = true;

        const typingDiv = document.createElement('div');
        typingDiv.className = 'message assistant';
        typingDiv.id = 'typingMessage';
        typingDiv.innerHTML = `
            <div class="message-header">
                <span>🤖 合规助手</span>
                <span>${new Date().toLocaleTimeString()}</span>
            </div>
            <div class="typing-indicator">
                <span></span>
                <span></span>
                <span></span>
                <small style="margin-left: 10px; color: #6c757d;">正在搜索知识库...</small>
            </div>
        `;

        this.messagesContainer.appendChild(typingDiv);
        this.scrollToBottom();
    }

    removeTyping() {
        const typingMessage = document.getElementById('typingMessage');
        if (typingMessage) {
            typingMessage.remove();
        }
        this.isTyping = false;
    }

    async sendMessage() {
        const message = this.userInput.value.trim();
        if (!message || this.isTyping) return;

        // 添加用户消息
        this.addMessage(message, 'user');
        this.userInput.value = '';
        this.sendBtn.disabled = true;

        // 显示输入状态
        this.showTyping();

        try {
            // 调用API
            const response = await fetch('/api/ask', {
                method: 'POST',
                headers: {
                    'Content-Type': 'application/json',
                },
                body: JSON.stringify({ question: message })
            });

            const data = await response.json();

            // 移除输入状态
            this.removeTyping();

            // 添加回答
            if (data.success) {
                this.displayAnswers(data);
            } else {
                this.addMessage(data.message, 'assistant');
            }

        } catch (error) {
            this.removeTyping();
            this.addMessage('抱歉，服务暂时不可用。请稍后重试。', 'assistant');
        } finally {
            this.sendBtn.disabled = false;
            this.userInput.focus();
        }
    }

    displayAnswers(data) {
        const answersHtml = [];

        if (data.high_similarity && data.high_similarity.length > 0) {
            answersHtml.push('<div class="answers-section"><h4>🎯 精准匹配</h4>');

            data.high_similarity.forEach((answer, index) => {
                answersHtml.push(`
                    <div class="answer-card high">
                        <div class="answer-title high">
                            <div class="answer-title-icon">✓</div>
                            <span>答案 ${answer.id}</span>
                            <span class="similarity-badge high">${Math.round(answer.similarity * 100)}%</span>
                        </div>
                        <div>${answer.answer}</div>
                    </div>
                `);
            });

            answersHtml.push('</div>');
        }

        if (data.low_similarity && data.low_similarity.length > 0) {
            answersHtml.push('<div class="answers-section"><h4>🔍 相关参考</h4>');

            data.low_similarity.forEach((answer, index) => {
                answersHtml.push(`
                    <div class="answer-card medium">
                        <div class="answer-title medium">
                            <div class="answer-title-icon">💡</div>
                            <span>参考 ${answer.id}</span>
                            <span class="similarity-badge medium">${Math.round(answer.similarity * 100)}%</span>
                        </div>
                        <div>${answer.answer}</div>
                    </div>
                `);
            });

            answersHtml.push('</div>');
        }

        const finalHtml = answersHtml.join('');
        this.addMessage(finalHtml, 'assistant');

        // 如果有关键词，显示提示
        if (data.keywords && data.keywords.length > 0) {
            setTimeout(() => {
                const keywordsHtml = `
                    <div style="margin-top: 15px; padding: 10px; background: #f8f9fa; border-radius: 10px;">
                        <small style="color: #6c757d;">
                            🔍 关键词：${data.keywords.join('、')}
                        </small>
                    </div>
                `;
                const lastMessage = this.messagesContainer.lastElementChild;
                if (lastMessage && lastMessage.classList.contains('assistant')) {
                    lastMessage.querySelector('.message-content').insertAdjacentHTML('beforeend', keywordsHtml);
                }
            }, 500);
        }
    }

    scrollToBottom() {
        this.messagesContainer.scrollTop = this.messagesContainer.scrollHeight;
    }
}

// 初始化应用
document.addEventListener('DOMContentLoaded', () => {
    new ChatApp();
});
'''

    with open('static/js/app.js', 'w', encoding='utf-8') as f:
        f.write(js_content)

    # 创建HTML模板
    html_content = '''<!DOCTYPE html>
<html lang="zh-CN">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>美易在线答疑 - 美纳里尼合规报销智能助手</title>
    <link href="https://cdn.jsdelivr.net/npm/bootstrap@5.1.3/dist/css/bootstrap.min.css" rel="stylesheet">
    <link rel="stylesheet" href="{{ url_for('static', filename='css/style.css') }}">
    <link rel="stylesheet" href="https://cdnjs.cloudflare.com/ajax/libs/font-awesome/6.0.0/css/all.min.css">
    <style>
        /* 配置面板样式 */
        .config-btn {
            position: fixed;
            top: 80px;
            right: 20px;
            width: 50px;
            height: 50px;
            background: rgba(255, 255, 255, 0.95);
            backdrop-filter: blur(10px);
            border-radius: 50%;
            border: none;
            box-shadow: 0 5px 15px rgba(0,0,0,0.1);
            cursor: pointer;
            display: flex;
            align-items: center;
            justify-content: center;
            font-size: 20px;
            color: var(--primary-color);
            transition: var(--transition);
            z-index: 1000;
        }

        .config-btn:hover {
            transform: rotate(90deg);
            box-shadow: 0 8px 25px rgba(0,0,0,0.15);
        }

        .config-panel, .upload-panel {
            position: fixed;
            top: 50%;
            left: 50%;
            transform: translate(-50%, -50%);
            background: white;
            border-radius: var(--border-radius);
            box-shadow: 0 20px 60px rgba(0,0,0,0.3);
            z-index: 1001;
            min-width: 400px;
            max-width: 500px;
        }

        .config-panel-header, .upload-panel-header {
            padding: 20px;
            border-bottom: 1px solid #e9ecef;
            display: flex;
            justify-content: space-between;
            align-items: center;
        }

        .config-panel-body, .upload-panel-body {
            padding: 20px;
        }

        .action-buttons {
            position: fixed;
            bottom: 20px;
            right: 20px;
            display: flex;
            gap: 10px;
            z-index: 1000;
        }

        .action-buttons .btn {
            border-radius: 25px;
            padding: 10px 20px;
            box-shadow: 0 5px 15px rgba(0,0,0,0.1);
            transition: var(--transition);
        }

        .action-buttons .btn:hover {
            transform: translateY(-2px);
            box-shadow: 0 8px 25px rgba(0,0,0,0.15);
        }
    </style>
</head>
<body>
    <!-- 状态指示器 -->
    <div class="status-badge offline" id="statusBadge">
        <span class="status-dot"></span>
        <span>检测中...</span>
    </div>

    <!-- 主容器 -->
    <div class="chat-container">
        <!-- 头部 -->
        <header class="header">
            <div class="header-icon">
                <i class="fas fa-robot"></i>
            </div>
            <div class="header-content">
                <h1>美易在线答疑</h1>
                <p>美纳里尼合规报销智能助手</p>
            </div>
        </header>

        <!-- 消息区域 -->
        <div class="messages-container">
            <div class="messages" id="messages">
                <!-- 初始欢迎消息 -->
                <div class="message assistant">
                    <div class="message-header">
                        <span>🤖 合规助手</span>
                        <span id="welcomeTime"></span>
                    </div>
                    <div class="message-content">
                        <strong>欢迎使用美纳里尼合规报销智能助手！</strong><br><br>
                        我可以帮您解答合规和报销相关问题。请输入您的问题，比如：
                        <br>• 为什么重置密码的邮件一直收不到？
                        <br>• 报销流程是什么？
                        <br>• 如何登录系统？
                        <br><br>
                        <div class="quick-questions">
                            <button class="quick-question">为什么重置密码的邮件一直收不到?</button>
                            <button class="quick-question">报销流程是什么?</button>
                            <button class="quick-question">如何登录系统?</button>
                            <button class="quick-question">发票要求有哪些?</button>
                        </div>
                    </div>
                </div>
            </div>
        </div>

        <!-- 输入区域 -->
        <div class="input-container">
            <form id="chatForm" class="input-group">
                <input type="text" class="form-control" id="userInput"
                       placeholder="请输入您的问题..." required>
                <button type="submit" class="btn btn-primary" id="sendBtn">
                    <i class="fas fa-paper-plane me-2"></i>发送
                </button>
            </form>
        </div>
    </div>

    <script src="https://cdn.jsdelivr.net/npm/bootstrap@5.1.3/dist/js/bootstrap.bundle.min.js"></script>
    <script src="{{ url_for('static', filename='js/app.js') }}"></script>
    <script>
        // 设置欢迎消息的时间
        document.getElementById('welcomeTime').textContent = new Date().toLocaleTimeString();
    </script>
</body>
</html>'''

    with open('templates/index_v2.html', 'w', encoding='utf-8') as f:
        f.write(html_content)

def main():
    """主函数"""
    # 创建模板和静态文件
    create_files()

    # 加载知识库
    load_knowledge_base()

    import argparse

    parser = argparse.ArgumentParser(description='美纳里尼合规报销智能助手Web应用 v2.0')
    parser.add_argument('--host', default='0.0.0.0', help='主机地址 (默认: 0.0.0.0)')
    parser.add_argument('--port', type=int, default=5000, help='端口号 (默认: 5000)')
    parser.add_argument('--debug', action='store_true', help='调试模式')

    args = parser.parse_args()

    print("Starting web application v2.0...")
    print("Local access: http://localhost:%d" % args.port)
    print("External access: http://%s:%d" % (args.host, args.port))
    print("Knowledge base Q&A pairs: %d" % len(qa_pairs))
    print("Press Ctrl+C to stop server")

    app.run(host=args.host, port=args.port, debug=args.debug)

if __name__ == "__main__":
    main()