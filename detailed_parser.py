#!/usr/bin/env python3
"""
详细解析Word文档
"""

import os
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

def analyze_document_structure():
    """分析文档结构"""
    print("开始分析文档结构...")

    rag_dir = r'C:\Users\jj\Desktop\RAG'
    doc_path = os.path.join(rag_dir, 'QA.docx')

    doc = Document(doc_path)

    print(f"文档统计:")
    print(f"- 总段落数: {len(doc.paragraphs)}")
    print(f"- 总表格数: {len(doc.tables)}")
    print(f"- 总节节数: {len(doc.sections)}")

    # 分析段落
    print(f"\n段落分析:")
    question_paragraphs = []
    for i, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if text and ('？' in text or '?' in text) and len(text) > 5:
            question_paragraphs.append((i, text))
            print(f"  可能的问题段落[{i+1}]: {text[:80]}...")

    print(f"\n找到 {len(question_paragraphs)} 个可能的问题段落")

    # 分析表格
    print(f"\n表格分析:")
    for i, table in enumerate(doc.tables):
        print(f"\n表格 {i+1}:")
        print(f"  行数: {len(table.rows)}")
        print(f"  列数: {len(table.columns)}")

        # 尝试不同的行和列组合
        for row_idx, row in enumerate(table.rows):
            row_text = []
            for cell_idx, cell in enumerate(row.cells):
                cell_text = cell.text.strip()
                row_text.append(f"列{cell_idx+1}: '{cell_text}'")

            print(f"  行{row_idx+1}: {' | '.join(row_text)}")

            # 检查是否包含问答
            for cell in row.cells:
                text = cell.text.strip()
                if '？' in text or '?' in text:
                    print(f"    -> 找到可能的问题: {text[:50]}...")

    # 尝试从段落中提取QA对
    print(f"\n尝试从段落中提取QA对...")
    qa_pairs = []

    # 方法1: 查找问句后跟答案的模式
    current_question = None
    for i, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if not text:
            continue

        # 如果是问句
        if ('？' in text or '?' in text) and len(text) > 5:
            # 保存上一个问题
            if current_question:
                # 找到下一个非问题段落作为答案
                for j in range(i+1, min(i+5, len(doc.paragraphs))):
                    next_para = doc.paragraphs[j]
                    next_text = next_para.text.strip()
                    if next_text and not ('？' in next_text or '?' in next_text):
                        qa_pairs.append({
                            'question': current_question,
                            'answer': next_text,
                            'method': '段落匹配'
                        })
                        break
            # 开始新问题
            current_question = text
        # 如果不是问句但有当前问题
        elif current_question and text:
            # 作为答案
            qa_pairs.append({
                'question': current_question,
                'answer': text,
                'method': '段落匹配'
            })
            current_question = None

    print(f"\n从段落提取了 {len(qa_pairs)} 个QA对")
    for i, qa in enumerate(qa_pairs[:10]):
        print(f"{i+1}. [{qa['method']}]")
        print(f"   问题: {qa['question']}")
        print(f"   答案: {qa['answer']}")

    # 搜索重置密码相关
    print(f"\n搜索重置密码相关QA:")
    password_related = []
    for qa in qa_pairs:
        if '重置密码' in qa['question'] or '重置密码' in qa['answer'] or '密码' in qa['question'] or '密码' in qa['answer']:
            password_related.append(qa)

    print(f"找到 {len(password_related)} 个密码相关的QA:")
    for i, qa in enumerate(password_related):
        print(f"\n{i+1}. [{qa['method']}]")
        print(f"   问题: {qa['question']}")
        print(f"   答案: {qa['answer']}")

    return qa_pairs

if __name__ == "__main__":
    qa_pairs = analyze_document_structure()

    # 保存结果
    if qa_pairs:
        import json
        with open("extracted_qa_pairs.json", "w", encoding="utf-8") as f:
            json.dump(qa_pairs, f, ensure_ascii=False, indent=2)
        print(f"\n结果已保存到 extracted_qa_pairs.json")