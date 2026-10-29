#!/usr/bin/env python3
"""
添加高级功能到应用
"""

import os
import json
import re
from typing import List, Dict, Any, Tuple
from datetime import datetime

class AdvancedSearch:
    """高级搜索功能"""

    def __init__(self):
        pass

    def parse_word_with_categories(self, file_path: str) -> Tuple[List[Dict[str, Any]], Dict[str, List[Dict[str, Any]]]]:
        """解析Word文档并分类"""
        try:
            from docx import Document
            doc = Document(file_path)
            qa_pairs = []
            categories = {
                '密码相关': [],
                '报销相关': [],
                '登录相关': [],
                '系统问题': [],
                '其他': []
            }

            # 从段落提取QA对
            for i in range(len(doc.paragraphs) - 1):
                para1 = doc.paragraphs[i]
                para2 = doc.paragraphs[i + 1]

                text1 = para1.text.strip()
                text2 = para2.text.strip()

                # 检查第一个段落是否是问题
                if text1 and ('？' in text1 or '?' in text1) and len(text1) > 5:
                    # 第二个段落作为答案
                    if text2 and len(text2) > 5:
                        qa = {
                            'id': len(qa_pairs) + 1,
                            'question': text1,
                            'answer': text2,
                            'source': os.path.basename(file_path),
                            'index': i + 1
                        }
                        qa_pairs.append(qa)

                        # 分类
                        if any(keyword in text1.lower() or keyword in text2.lower()
                               for keyword in ['密码', '重置', '邮箱']):
                            categories['密码相关'].append(qa)
                        elif any(keyword in text1.lower() or keyword in text2.lower()
                                for keyword in ['报销', '发票', '财务', '申请']):
                            categories['报销相关'].append(qa)
                        elif any(keyword in text1.lower() or keyword in text2.lower()
                                for keyword in ['登录', '账号', '账户']):
                            categories['登录相关'].append(qa)
                        elif any(keyword in text1.lower() or keyword in text2.lower()
                                for keyword in ['系统', '平台', '网站']):
                            categories['系统问题'].append(qa)
                        else:
                            categories['其他'].append(qa)

            return qa_pairs, categories
        except Exception as e:
            print(f"解析失败: {e}")
            return [], {}

    def search_with_fuzzy(self, query: str, qa_pairs: List[Dict[str, Any]],
                         max_answers: int = 5) -> List[Dict[str, Any]]:
        """模糊搜索，支持同义词扩展"""
        query_lower = query.lower()
        results = []

        # 定义同义词词典
        synonyms = {
            '重置密码': ['重置密码', '重置口令', '找回密码', '修改密码'],
            '邮箱': ['邮箱', '电子邮箱', 'email', '邮件'],
            '收不到': ['收不到', '未收到', '没有收到', '没收到'],
            '登录': ['登录', '登入', ' signin', 'log in'],
            '报销': ['报销', '报销流程', '报销申请', '报销系统']
        }

        # 扩展查询词
        expanded_terms = set()
        for term, synonyms_list in synonyms.items():
            if term in query_lower:
                expanded_terms.update(synonyms_list)

        # 如果没有找到同义词，使用原查询
        if not expanded_terms:
            expanded_terms.add(query_lower)

        for qa in qa_pairs:
            score = 0
            q_text = qa['question'].lower()
            a_text = qa['answer'].lower()

            # 计算每个扩展词的匹配度
            for term in expanded_terms:
                if term in q_text:
                    score += 1.0
                elif term in a_text:
                    score += 0.7
                else:
                    # 部分匹配
                    words = term.split()
                    match_count = sum(1 for word in words if word in q_text or word in a_text)
                    if match_count > 0:
                        score += match_count / len(words) * 0.5

            if score > 0:
                results.append({
                    'id': qa['id'],
                    'question': qa['question'],
                    'answer': qa['answer'],
                    'source': qa['source'],
                    'similarity': score / len(expanded_terms),
                    'index': qa.get('index', 0)
                })

        # 按相似度排序
        results.sort(key=lambda x: x['similarity'], reverse=True)
        return results[:max_answers]

    def get_suggested_questions(self, keywords: List[str], qa_pairs: List[Dict[str, Any]],
                               max_suggestions: int = 5) -> List[str]:
        """根据关键词获取建议问题"""
        suggestions = set()

        for qa in qa_pairs:
            for keyword in keywords:
                if keyword.lower() in qa['question'].lower():
                    suggestions.add(qa['question'])
                    if len(suggestions) >= max_suggestions:
                        return list(suggestions)

        return list(suggestions)

def create_enhanced_app():
    """创建增强版应用"""

    # 1. 准备知识库数据
    print("准备知识库数据...")
    search_engine = AdvancedSearch()

    rag_dir = r'C:\Users\jj\Desktop\RAG'
    qa_pairs = []
    categories = {}

    for file in os.listdir(rag_dir):
        if file == 'QA.docx':
            file_path = os.path.join(rag_dir, file)
            qa_pairs, categories = search_engine.parse_word_with_categories(file_path)
            print(f"加载了 {len(qa_pairs)} 个QA对")
            break

    # 2. 保存分类数据
    with open("knowledge_base_categories.json", "w", encoding="utf-8") as f:
        json.dump(categories, f, ensure_ascii=False, indent=2)

    # 3. 创建推荐问题
    print("生成推荐问题...")
    common_questions = {
        '密码相关': ['重置密码', '登录问题', '邮箱验证'],
        '报销相关': ['报销流程', '发票要求', '报销周期'],
        '登录相关': ['无法登录', '账号问题', '密码错误'],
        '系统问题': ['系统使用', '平台功能', '网站问题']
    }

    suggestions = {}
    for category, keywords in common_questions.items():
        suggestions[category] = search_engine.get_suggested_questions(
            keywords, qa_pairs, max_suggestions=3
        )

    # 4. 保存建议数据
    data = {
        'qa_pairs': qa_pairs,
        'categories': categories,
        'suggestions': suggestions,
        'last_updated': datetime.now().isoformat()
    }

    with open("enhanced_knowledge_base.json", "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)

    print("增强版知识库已准备完成！")

    # 5. 显示统计信息
    print("\n知识库统计:")
    print(f"- 总问答对: {len(qa_pairs)}")
    for category, qas in categories.items():
        print(f"- {category}: {len(qas)} 个")

    print("\n建议问题:")
    for category, qs in suggestions.items():
        print(f"- {category}:")
        for q in qs:
            print(f"  * {q}")

if __name__ == "__main__":
    create_enhanced_app()