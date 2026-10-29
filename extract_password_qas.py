#!/usr/bin/env python3
"""
提取所有密码相关的问题
"""

import os
from docx import Document

def extract_all_qa_pairs():
    """提取所有QA对"""
    print("开始提取所有QA对...")

    rag_dir = r'C:\Users\jj\Desktop\RAG'
    doc_path = os.path.join(rag_dir, 'QA.docx')

    doc = Document(doc_path)
    qa_pairs = []

    # 从段落提取QA对
    print("\n从段落提取QA对...")
    for i in range(len(doc.paragraphs) - 1):
        para1 = doc.paragraphs[i]
        para2 = doc.paragraphs[i + 1]

        text1 = para1.text.strip()
        text2 = para2.text.strip()

        # 检查第一个段落是否是问题
        if text1 and ('？' in text1 or '?' in text1) and len(text1) > 5:
            # 第二个段落作为答案
            if text2 and len(text2) > 5:
                qa_pairs.append({
                    'question': text1,
                    'answer': text2,
                    'index': i + 1
                })

    print(f"共提取了 {len(qa_pairs)} 个QA对")
    return qa_pairs

def find_password_related(qa_pairs):
    """查找密码相关的问题"""
    print("\n查找密码相关的问题...")

    password_keywords = ['密码', '重置', '登录', '账号', '账户', '验证', '邮箱']
    related_qas = []

    for qa in qa_pairs:
        # 检查问题或答案中是否包含关键词
        question_lower = qa['question'].lower()
        answer_lower = qa['answer'].lower()

        has_keyword = any(keyword in question_lower or keyword in answer_lower
                         for keyword in password_keywords)

        if has_keyword:
            related_qas.append(qa)

    print(f"找到 {len(related_qas)} 个密码相关的QA:")

    for i, qa in enumerate(related_qas, 1):
        print(f"\n{i}. [位置: {qa['index']}]")
        print(f"   问题: {qa['question']}")
        print(f"   答案: {qa['answer']}")
        print()

    return related_qas

def find_specific_password_questions(qa_pairs):
    """查找特定密码问题"""
    print("\n查找特定密码问题...")

    # 查找包含"重置密码"的问题
    reset_password_qas = []
    for qa in qa_pairs:
        if '重置密码' in qa['question']:
            reset_password_qas.append(qa)

    print(f"包含'重置密码'的问题（共 {len(reset_password_qas)} 个）:")
    for i, qa in enumerate(reset_password_qas, 1):
        print(f"\n{i}. [位置: {qa['index']}]")
        print(f"   问题: {qa['question']}")
        print(f"   答案: {qa['answer']}")

    # 查找包含"邮箱"的问题
    email_qas = []
    for qa in qa_pairs:
        if '邮箱' in qa['question'] or 'email' in qa['question'].lower():
            email_qas.append(qa)

    print(f"\n包含'邮箱'的问题（共 {len(email_qas)} 个）:")
    for i, qa in enumerate(email_qas, 1):
        print(f"\n{i}. [位置: {qa['index']}]")
        print(f"   问题: {qa['question']}")
        print(f"   答案: {qa['answer']}")

    # 查找包含"收不到"的问题
    no_receive_qas = []
    for qa in qa_pairs:
        if '收不到' in qa['question'] or '未收到' in qa['question']:
            no_receive_qas.append(qa)

    print(f"\n包含'收不到'的问题（共 {len(no_receive_qas)} 个）:")
    for i, qa in enumerate(no_receive_qas, 1):
        print(f"\n{i}. [位置: {qa['index']}]")
        print(f"   问题: {qa['question']}")
        print(f"   答案: {qa['answer']}")

    return reset_password_qas, email_qas, no_receive_qas

def get_all_answers_for_user_question(user_question, qa_pairs):
    """获取用户问题的所有相关答案"""
    print(f"\n为用户问题查找所有相关答案: {user_question}")

    user_keywords = user_question.lower().split()
    related_qas = []

    for qa in qa_pairs:
        score = 0
        question_lower = qa['question'].lower()
        answer_lower = qa['answer'].lower()

        # 计算匹配度
        for keyword in user_keywords:
            if keyword in question_lower:
                score += 1
            if keyword in answer_lower:
                score += 0.5

        if score > 0:
            related_qas.append({
                'qa': qa,
                'score': score
            })

    # 按分数排序
    related_qas.sort(key=lambda x: x['score'], reverse=True)

    print(f"找到 {len(related_qas)} 个相关答案:")
    for i, item in enumerate(related_qas[:5], 1):
        qa = item['qa']
        score = item['score']
        print(f"\n{i}. [匹配度: {score}] [位置: {qa['index']}]")
        print(f"   问题: {qa['question']}")
        print(f"   答案: {qa['answer']}")

    return related_qas

def main():
    """主函数"""
    # 1. 提取所有QA对
    qa_pairs = extract_all_qa_pairs()

    # 2. 查找密码相关
    password_qas = find_password_related(qa_pairs)

    # 3. 查找特定密码问题
    reset_qas, email_qas, no_receive_qas = find_specific_password_questions(qa_pairs)

    # 4. 查找用户问题的所有答案
    user_question = "为什么重置密码的邮件一直收不到?"
    related_answers = get_all_answers_for_user_question(user_question, qa_pairs)

    # 5. 保存结果
    results = {
        'total_qa_pairs': len(qa_pairs),
        'password_related': len(password_qas),
        'reset_password_questions': reset_qas,
        'email_questions': email_qas,
        'no_receive_questions': no_receive_qas,
        'user_question_related': related_answers
    }

    import json
    with open("password_qa_results.json", "w", encoding="utf-8") as f:
        json.dump(results, f, ensure_ascii=False, indent=2)

    print(f"\n结果已保存到 password_qa_results.json")

    # 6. 总结
    print(f"\n总结:")
    print(f"- 文档中共有 {len(qa_pairs)} 个QA对")
    print(f"- 其中 {len(password_qas)} 个与密码相关")
    print(f"- 找到 {len(reset_qas)} 个关于重置密码的问题")
    print(f"- 找到 {len(email_qas)} 个关于邮箱的问题")
    print(f"- 找到 {len(no_receive_qas)} 个关于收不到的问题")
    print(f"- 用户问题找到 {len(related_answers)} 个相关答案")

if __name__ == "__main__":
    main()