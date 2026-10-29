#!/usr/bin/env python3
"""
美纳里尼合规报销智能助手 - 最终优化版
包含所有原版功能
"""

from flask import Flask, request, jsonify, render_template, send_from_directory
import os
import re
import werkzeug.utils
from typing import List, Dict, Any
from datetime import datetime

app = Flask(__name__)

# 配置参数
MAX_ANSWERS = 2
SIMILARITY_THRESHOLD = 0.4

# 知识库数据
qa_pairs = []
initialized = False

class SearchEngine:
    """搜索引擎"""

    def __init__(self):
        pass

    def parse_qa_pairs(self) -> List[Dict[str, Any]]:
        """解析QA对"""
        try:
            from docx import Document
            doc = Document("C:\\\\Users\\\\jj\\\\Desktop\\\\RAG\\\\QA.docx")
            qa_pairs = []

            for i in range(len(doc.paragraphs) - 1):
                para1 = doc.paragraphs[i]
                para2 = doc.paragraphs[i + 1]

                text1 = para1.text.strip()
                text2 = para2.text.strip()

                if text1 and ('？' in text1 or '?' in text1) and len(text1) > 5:
                    if text2 and len(text2) > 5:
                        qa = {
                            'id': len(qa_pairs) + 1,
                            'question': text1,
                            'answer': text2,
                            'index': i + 1
                        }
                        qa_pairs.append(qa)

            return qa_pairs
        except Exception as e:
            print(f"解析失败: {e}")
            return []

    def search_with_fuzzy(self, query: str, qa_pairs: List[Dict[str, Any]], max_answers: int = MAX_ANSWERS) -> List[Dict[str, Any]]:
        """模糊搜索"""
        query_lower = query.lower()
        results = []

        # 同义词词典
        synonyms = {
            '重置密码': ['重置密码', '重置口令', '找回密码', '修改密码'],
            '邮箱': ['邮箱', '电子邮箱', 'email', '邮件'],
            '收不到': ['收不到', '未收到', '没有收到', '没收到'],
            '登录': ['登录', '登入', 'signin', 'log in'],
            '报销': ['报销', '报销流程', '报销申请', '报销系统']
        }

        # 扩展查询词
        expanded_terms = set()
        for term, synonyms_list in synonyms.items():
            if term in query_lower:
                expanded_terms.update(synonyms_list)

        if not expanded_terms:
            expanded_terms.add(query_lower)

        # 计算相似度
        for qa in qa_pairs:
            score = 0
            q_text = qa['question'].lower()
            a_text = qa['answer'].lower()

            for term in expanded_terms:
                if term in q_text:
                    score += 1.0
                elif term in a_text:
                    score += 0.7
                else:
                    words = term.split()
                    match_count = sum(1 for word in words if word in q_text or word in a_text)
                    if match_count > 0:
                        score += match_count / len(words) * 0.5

            if score > 0:
                results.append({
                    'id': qa['id'],
                    'question': qa['question'],
                    'answer': qa['answer'],
                    'similarity': score / len(expanded_terms),
                    'index': qa.get('index', 0)
                })

        # 按相似度排序
        results.sort(key=lambda x: x['similarity'], reverse=True)
        return results[:max_answers]

    def highlight_keywords(self, text: str, keywords: List[str]) -> str:
        """高亮关键词"""
        if not keywords:
            return text

        for keyword in keywords:
            if keyword:
                text = re.sub(
                    re.escape(keyword),
                    f'<mark class="highlight-mark">{keyword}</mark>',
                    text,
                    flags=re.IGNORECASE
                )
        return text

    def extract_keywords(self, text: str) -> List[str]:
        """提取关键词"""
        words = re.findall(r'[\\u4e00-\\u9fff]+|[a-zA-Z]+', text)
        stopwords = {'的', '了', '在', '是', '我', '有', '和', '就', '不', '人', '都', '一', '个', '上', '也', '很', '到', '说', '要', '去', '你', '会', '着', '没有', '看', '好', '自己', '这'}
        keywords = []
        for word in words:
            if len(word) >= 2 and word not in stopwords and word not in keywords:
                keywords.append(word)
        return keywords[:5]

# 全局搜索引擎实例
search_engine = SearchEngine()

def load_knowledge_base():
    """加载知识库"""
    global qa_pairs, initialized
    if not initialized:
        qa_pairs = search_engine.parse_qa_pairs()
        initialized = True
        print(f"知识库加载成功，共 {len(qa_pairs)} 个问答对")

# API路由
@app.route('/')
def index():
    """主页"""
    return render_template('index_final.html')

@app.route('/api/ask', methods=['POST'])
def ask():
    """问答API"""
    try:
        data = request.json
        question = data.get('question', '').strip()

        if not question:
            return jsonify({'error': '请输入问题'}), 400

        # 获取配置参数
        max_answers = data.get('maxAnswers', MAX_ANSWERS)
        similarity_threshold = data.get('similarityThreshold', SIMILARITY_THRESHOLD)

        # 加载知识库
        if not initialized:
            load_knowledge_base()

        if not initialized:
            return jsonify({'error': '知识库未加载'}), 500

        # 搜索答案
        results = search_engine.search_with_fuzzy(question, qa_pairs, max_answers=max_answers)

        if results:
            # 提取关键词并高亮
            keywords = search_engine.extract_keywords(question)

            # 按相似度分组
            high_similarity = [r for r in results if r['similarity'] >= similarity_threshold]
            low_similarity = [r for r in results if r['similarity'] < similarity_threshold and r['similarity'] > 0]

            # 构建回答
            response_data = {
                'success': True,
                'high_similarity': [],
                'low_similarity': [],
                'keywords': keywords
            }

            for result in high_similarity:
                answer = search_engine.highlight_keywords(result['answer'], keywords)
                response_data['high_similarity'].append({
                    'id': result['id'],
                    'answer': answer,
                    'similarity': result['similarity'],
                    'original_question': result['question']
                })

            for result in low_similarity:
                answer = search_engine.highlight_keywords(result['answer'], keywords)
                response_data['low_similarity'].append({
                    'id': result['id'],
                    'answer': answer,
                    'similarity': result['similarity'],
                    'original_question': result['question']
                })

            return jsonify(response_data)
        else:
            return jsonify({
                'success': False,
                'error': '未找到相关答案',
                'message': '''很抱歉，我在知识库中没有找到与您问题相关的信息。

📝 建议：
1. 请检查问题表述是否清晰
2. 尝试使用不同的关键词描述您的问题
3. 如果仍需帮助，请联系人工客服

📞 客服联系方式：请通过公司内部通讯系统联系IT支持团队'''
            })

    except Exception as e:
        return jsonify({'error': f'服务器错误: {str(e)}'}), 500

@app.route('/api/config', methods=['GET', 'POST'])
def config():
    """配置API"""
    global MAX_ANSWERS, SIMILARITY_THRESHOLD
    if request.method == 'GET':
        return jsonify({
            'maxAnswers': MAX_ANSWERS,
            'similarityThreshold': SIMILARITY_THRESHOLD
        })
    else:
        try:
            data = request.json
            if 'maxAnswers' in data:
                MAX_ANSWERS = int(data['maxAnswers'])
            if 'similarityThreshold' in data:
                SIMILARITY_THRESHOLD = float(data['similarityThreshold'])
            return jsonify({'success': True})
        except Exception as e:
            return jsonify({'error': str(e)}), 400

@app.route('/api/upload', methods=['POST'])
def upload_knowledge_base():
    """上传知识库"""
    try:
        if 'file' not in request.files:
            return jsonify({'error': '没有文件上传'}), 400

        file = request.files['file']
        if file.filename == '':
            return jsonify({'error': '没有选择文件'}), 400

        if file and file.filename.endswith('.docx'):
            # 保存上传的文件
            upload_path = os.path.join('uploads', file.filename)
            os.makedirs('uploads', exist_ok=True)
            file.save(upload_path)

            # 重新解析知识库
            global qa_pairs, initialized
            qa_pairs = []
            initialized = False

            # 解析新上传的文件
            try:
                from docx import Document
                doc = Document(upload_path)
                qa_pairs = []

                for i in range(len(doc.paragraphs) - 1):
                    para1 = doc.paragraphs[i]
                    para2 = doc.paragraphs[i + 1]

                    text1 = para1.text.strip()
                    text2 = para2.text.strip()

                    if text1 and ('？' in text1 or '?' in text1) and len(text1) > 5:
                        if text2 and len(text2) > 5:
                            qa = {
                                'id': len(qa_pairs) + 1,
                                'question': text1,
                                'answer': text2,
                                'index': i + 1
                            }
                            qa_pairs.append(qa)

                initialized = True
                return jsonify({
                    'success': True,
                    'message': '知识库更新成功',
                    'qa_pairs_count': len(qa_pairs)
                })
            except Exception as e:
                return jsonify({'error': f'解析失败: {str(e)}'}), 500
        else:
            return jsonify({'error': '只支持.docx文件'}), 400
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/status')
def status():
    """状态API"""
    load_knowledge_base()
    return jsonify({
        'status': 'ok',
        'initialized': initialized,
        'qa_pairs_count': len(qa_pairs),
        'timestamp': datetime.now().isoformat()
    })

@app.route('/static/<path:filename>')
def static_files(filename):
    """静态文件"""
    return send_from_directory('static', filename)

def main():
    """主函数"""
    import argparse

    parser = argparse.ArgumentParser(description='美纳里尼合规报销智能助手')
    parser.add_argument('--host', default='0.0.0.0', help='主机地址 (默认: 0.0.0.0)')
    parser.add_argument('--port', type=int, default=5000, help='端口号 (默认: 5000)')

    args = parser.parse_args()

    print("Starting web application...")
    print("Local access: http://localhost:%d" % args.port)
    print("External access: http://%s:%d" % (args.host, args.port))
    print("Knowledge base Q&A pairs: %d" % len(qa_pairs))
    print("Press Ctrl+C to stop server")

    app.run(host=args.host, port=args.port)

if __name__ == "__main__":
    main()