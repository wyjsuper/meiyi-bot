import os
import openpyxl
from typing import List, Dict, Any, Optional
import logging

class RobustParser:
    """健壮的文档解析器，能够处理各种格式的Excel文件"""

    def __init__(self):
        self.logger = logging.getLogger(__name__)

    def parse_excel_safe(self, file_path: str) -> List[Dict[str, Any]]:
        """
        安全解析Excel文件，尝试多种方法

        Args:
            file_path: Excel文件路径

        Returns:
            问答对列表
        """
        qa_pairs = []

        try:
            # 方法1: 使用openpyxl正常读取
            try:
                wb = openpyxl.load_workbook(file_path, data_only=True)
                qa_pairs = self._extract_qa_pairs(wb, file_path)
                if qa_pairs:
                    self.logger.info(f"方法1成功，提取 {len(qa_pairs)} 个问答对")
                    return qa_pairs
            except Exception as e:
                self.logger.warning(f"方法1失败: {e}")

            # 方法2: 尝试读取所有工作表
            try:
                wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
                for sheet_name in wb.sheetnames:
                    try:
                        sheet = wb[sheet_name]
                        pairs = self._extract_qa_pairs_from_sheet(sheet, file_path, sheet_name)
                        qa_pairs.extend(pairs)
                    except Exception as e:
                        self.logger.warning(f"工作表 {sheet_name} 解析失败: {e}")
                        continue

                if qa_pairs:
                    self.logger.info(f"方法2成功，提取 {len(qa_pairs)} 个问答对")
                    return qa_pairs
            except Exception as e:
                self.logger.warning(f"方法2失败: {e}")

            # 方法3: 尝试逐行读取
            try:
                wb = openpyxl.load_workbook(file_path, data_only=True)
                for sheet_name in wb.sheetnames:
                    sheet = wb[sheet_name]
                    pairs = self._extract_qa_pairs_simple(sheet, file_path, sheet_name)
                    qa_pairs.extend(pairs)

                if qa_pairs:
                    self.logger.info(f"方法3成功，提取 {len(qa_pairs)} 个问答对")
                    return qa_pairs
            except Exception as e:
                self.logger.warning(f"方法3失败: {e}")

        except Exception as e:
            self.logger.error(f"所有方法都失败了: {e}")

        return qa_pairs

    def _extract_qa_pairs(self, wb, file_path: str) -> List[Dict[str, Any]]:
        """从工作簿中提取问答对"""
        qa_pairs = []

        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            pairs = self._extract_qa_pairs_from_sheet(sheet, file_path, sheet_name)
            qa_pairs.extend(pairs)

        return qa_pairs

    def _extract_qa_pairs_from_sheet(self, sheet, file_path: str, sheet_name: str) -> List[Dict[str, Any]]:
        """从工作表中提取问答对"""
        qa_pairs = []

        try:
            # 获取第一行作为可能的表头
            first_row = list(sheet.iter_rows(min_row=1, max_row=1, values_only=True))[0]

            # 尝试识别问题列和答案列
            question_col = None
            answer_col = None

            for i, header in enumerate(first_row):
                header_str = str(header).strip() if header else ""
                if '问题' in header_str or 'question' in header_str.lower():
                    question_col = i
                elif '答案' in header_str or 'answer' in header_str.lower():
                    answer_col = i

            # 如果找到了问题列和答案列
            if question_col is not None and answer_col is not None:
                for row_idx, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
                    if len(row) > max(question_col, answer_col):
                        question = str(row[question_col]) if row[question_col] else ""
                        answer = str(row[answer_col]) if row[answer_col] else ""

                        if question and answer:
                            qa_pairs.append({
                                'id': len(qa_pairs) + 1,
                                'question': question,
                                'answer': answer,
                                'source': os.path.basename(file_path),
                                'metadata': {
                                    'sheet_name': sheet_name,
                                    'row_number': row_idx
                                }
                            })

        except Exception as e:
            self.logger.warning(f"工作表 {sheet_name} 提取失败: {e}")

        return qa_pairs

    def _extract_qa_pairs_simple(self, sheet, file_path: str, sheet_name: str) -> List[Dict[str, Any]]:
        """简单方法提取问答对"""
        qa_pairs = []

        try:
            # 检查前几行是否包含问答
            for row_idx, row in enumerate(sheet.iter_rows(min_row=1, max_row=50, values_only=True), 1):
                if len(row) >= 2:
                    # 假设第一列是问题，第二列是答案
                    question = str(row[0]) if row[0] else ""
                    answer = str(row[1]) if row[1] else ""

                    # 检查是否是合理的问答
                    if (question and len(question) > 2 and
                        answer and len(answer) > 2 and
                        ('？' in question or '?' in question)):
                        qa_pairs.append({
                            'id': len(qa_pairs) + 1,
                            'question': question,
                            'answer': answer,
                            'source': os.path.basename(file_path),
                            'metadata': {
                                'sheet_name': sheet_name,
                                'row_number': row_idx
                            }
                        })

                        # 限制数量
                        if len(qa_pairs) >= 100:
                            break

        except Exception as e:
            self.logger.warning(f"简单方法提取失败: {e}")

        return qa_pairs