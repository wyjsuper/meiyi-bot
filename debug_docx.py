#!/usr/bin/env python3
"""
调试Word文档解析
"""

import os
from docx import Document

def debug_docx_parsing():
    """调试Word文档解析"""
    print("开始调试Word文档解析...")

    rag_dir = r'C:\Users\jj\Desktop\RAG'
    doc_path = os.path.join(rag_dir, 'QA.docx')

    if not os.path.exists(doc_path):
        print(f"文件不存在: {doc_path}")
        return

    print(f"文件路径: {doc_path}")
    print(f"文件大小: {os.path.getsize(doc_path)} bytes")

    try:
        # 加载文档
        doc = Document(doc_path)
        print(f"\n文档基本信息:")
        print(f"- 段落数: {len(doc.paragraphs)}")
        print(f"- 表格数: {len(doc.tables)}")

        # 检查前几个段落
        print(f"\n前5个段落:")
        for i, para in enumerate(doc.paragraphs[:5]):
            print(f"段落{i+1}: {para.text[:100]}...")

        # 检查表格
        print(f"\n表格信息:")
        for i, table in enumerate(doc.tables):
            print(f"表格{i+1}: {len(table.rows)}行 x {len(table.columns)}列")

            # 显示前3行
            print(f"表格{i+1}前3行:")
            for j, row in enumerate(table.rows[:3]):
                row_text = " | ".join([cell.text.strip() for cell in row.cells])
                print(f"  行{j+1}: {row_text}")

        # 尝试提取所有包含问号的行
        print(f"\n尝试提取问答对...")
        qa_pairs = []

        # 从表格提取
        for table_idx, table in enumerate(doc.tables):
            print(f"\n处理表格 {table_idx+1}:")
            for row_idx, row in enumerate(table.rows):
                if len(row.cells) >= 2:
                    q = row.cells[0].text.strip()
                    a = row.cells[1].text.strip()

                    if q and a and ('？' in q or '?' in q):
                        qa_pairs.append({
                            'question': q,
                            'answer': a,
                            'table': table_idx + 1,
                            'row': row_idx + 1
                        })
                        print(f"  找到QA: {q[:50]}...")

        print(f"\n总共找到 {len(qa_pairs)} 个QA对")

        # 搜索重置密码相关
        print(f"\n搜索重置密码相关:")
        for qa in qa_pairs:
            if '重置密码' in qa['question'] or '重置密码' in qa['answer']:
                print(f"找到相关QA:")
                print(f"  问题: {qa['question']}")
                print(f"  答案: {qa['answer']}")
                print(f"  位置: 表格{qa['table']}-行{qa['row']}")

    except Exception as e:
        print(f"解析出错: {e}")
        import traceback
        traceback.print_exc()

if __name__ == "__main__":
    debug_docx_parsing()