#!/usr/bin/env python3
"""
美易机器人Web应用版本
使用Flask框架，支持外网访问
"""

from flask import Flask, render_template, request, jsonify, send_from_directory
import os
import json
import re
from typing import List, Dict, Any
import threading
import time
from datetime import datetime

app = Flask(__name__)

# 初始化知识库数据
qa_pairs = []
initialized = False

class SearchEngine:
    """搜索引擎"""

    def __init__(self):
        pass

    def parse_qa_pairs(self) -> List[Dict[str, Any]]:
        """解析QA对"""
        try:
            from docx import Document
            doc = Document("C:\\\\Users\\\\jj\\\\Desktop\\\\RAG\\\\QA.docx")
            qa_pairs = []

            for i in range(len(doc.paragraphs) - 1):
                para1 = doc.paragraphs[i]
                para2 = doc.paragraphs[i + 1]

                text1 = para1.text.strip()
                text2 = para2.text.strip()

                if text1 and ('？' in text1 or '?' in text1) and len(text1) > 5:
                    if text2 and len(text2) > 5:
                        qa = {
                            'id': len(qa_pairs) + 1,
                            'question': text1,
                            'answer': text2,
                            'index': i + 1
                        }
                        qa_pairs.append(qa)

            return qa_pairs
        except Exception as e:
            print(f"解析失败: {e}")
            return []

    def search_with_fuzzy(self, query: str, qa_pairs: List[Dict[str, Any]], max_answers: int = 5) -> List[Dict[str, Any]]:
        """模糊搜索"""
        query_lower = query.lower()
        results = []

        # 定义同义词词典
        synonyms = {
            '重置密码': ['重置密码', '重置口令', '找回密码', '修改密码'],
            '邮箱': ['邮箱', '电子邮箱', 'email', '邮件'],
            '收不到': ['收不到', '未收到', '没有收到', '没收到'],
            '登录': ['登录', '登入', 'signin', 'log in'],
            '报销': ['报销', '报销流程', '报销申请', '报销系统']
        }

        # 扩展查询词
        expanded_terms = set()
        for term, synonyms_list in synonyms.items():
            if term in query_lower:
                expanded_terms.update(synonyms_list)

        # 如果没有找到同义词，使用原查询
        if not expanded_terms:
            expanded_terms.add(query_lower)

        for qa in qa_pairs:
            score = 0
            q_text = qa['question'].lower()
            a_text = qa['answer'].lower()

            # 计算每个扩展词的匹配度
            for term in expanded_terms:
                if term in q_text:
                    score += 1.0
                elif term in a_text:
                    score += 0.7
                else:
                    # 部分匹配
                    words = term.split()
                    match_count = sum(1 for word in words if word in q_text or word in a_text)
                    if match_count > 0:
                        score += match_count / len(words) * 0.5

            if score > 0:
                results.append({
                    'id': qa['id'],
                    'question': qa['question'],
                    'answer': qa['answer'],
                    'similarity': score / len(expanded_terms),
                    'index': qa.get('index', 0)
                })

        # 按相似度排序
        results.sort(key=lambda x: x['similarity'], reverse=True)
        return results[:max_answers]

    def highlight_keywords(self, text: str, keywords: List[str]) -> str:
        """高亮关键词"""
        if not keywords:
            return text

        for keyword in keywords:
            if keyword:
                text = re.sub(
                    re.escape(keyword),
                    f'<mark style="background-color: #ffff00; padding: 2px 4px; border-radius: 3px;">{keyword}</mark>',
                    text,
                    flags=re.IGNORECASE
                )
        return text

    def extract_keywords(self, text: str) -> List[str]:
        """提取关键词"""
        words = re.findall(r'[\\u4e00-\\u9fff]+|[a-zA-Z]+', text)
        stopwords = {'的', '了', '在', '是', '我', '有', '和', '就', '不', '人', '都', '一', '个', '上', '也', '很', '到', '说', '要', '去', '你', '会', '着', '没有', '看', '好', '自己', '这'}
        keywords = []
        for word in words:
            if len(word) >= 2 and word not in stopwords and word not in keywords:
                keywords.append(word)
        return keywords[:5]

# 全局搜索引擎实例
search_engine = SearchEngine()

def load_knowledge_base():
    """加载知识库"""
    global qa_pairs, initialized
    if not initialized:
        try:
            qa_pairs = search_engine.parse_qa_pairs()
            initialized = True
            print(f"知识库加载成功，共 {len(qa_pairs)} 个问答对")
        except Exception as e:
            print(f"知识库加载失败: {e}")

# 路由定义
@app.route('/')
def index():
    """主页"""
    return render_template('index.html')

@app.route('/api/ask', methods=['POST'])
def ask():
    """问答API"""
    try:
        data = request.json
        question = data.get('question', '').strip()

        if not question:
            return jsonify({'error': '请输入问题'}), 400

        # 加载知识库
        if not initialized:
            load_knowledge_base()

        if not initialized:
            return jsonify({'error': '知识库未加载'}), 500

        # 搜索答案
        results = search_engine.search_with_fuzzy(question, qa_pairs, max_answers=5)

        if results:
            # 提取关键词并高亮
            keywords = search_engine.extract_keywords(question)
            answer_parts = []

            # 按相似度分组
            high_similarity = [r for r in results if r['similarity'] >= 0.2]
            low_similarity = [r for r in results if r['similarity'] < 0.2 and r['similarity'] > 0]

            if high_similarity:
                answer_parts.append("### 🎯 高相关答案")
                for i, result in enumerate(high_similarity, 1):
                    answer = search_engine.highlight_keywords(result['answer'], keywords)
                    answer_parts.append(f"**答案{i}**（相似度: {result['similarity']:.2%}）：\\n{answer}")

            if low_similarity and len(results) > len(high_similarity):
                answer_parts.append("\\n### 🔍 相关参考答案")
                for i, result in enumerate(low_similarity[:3], 1):
                    answer = search_engine.highlight_keywords(result['answer'], keywords)
                    answer_parts.append(f"**参考{i}**（相似度: {result['similarity']:.2%}）：\\n{answer}")

            answer = "\\n\\n".join(answer_parts)

            return jsonify({
                'success': True,
                'answer': answer,
                'results_count': len(results)
            })
        else:
            return jsonify({
                'success': False,
                'answer': """很抱歉，我在知识库中没有找到与您问题相关的信息。

📝 建议：
1. 请检查问题表述是否清晰
2. 尝试使用不同的关键词描述您的问题
3. 如果仍需帮助，请联系人工客服

📞 客服联系方式：请通过公司内部通讯系统联系IT支持团队""",
                'results_count': 0
            })

    except Exception as e:
        return jsonify({'error': f'服务器错误: {str(e)}'}), 500

@app.route('/api/status')
def status():
    """获取状态"""
    return jsonify({
        'status': 'ok',
        'initialized': initialized,
        'qa_pairs_count': len(qa_pairs),
        'timestamp': datetime.now().isoformat()
    })

@app.route('/static/<path:filename>')
def static_files(filename):
    """静态文件"""
    return send_from_directory('static', filename)

# 创建模板目录
def create_template():
    """创建模板文件"""
    os.makedirs('templates', exist_ok=True)

    index_html = '''<!DOCTYPE html>
<html lang="zh-CN">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>美易在线答疑 - 美纳里尼中国合规智能助手</title>
    <link href="https://cdn.jsdelivr.net/npm/bootstrap@5.1.3/dist/css/bootstrap.min.css" rel="stylesheet">
    <style>
        :root {
            --primary-color: #667eea;
            --secondary-color: #764ba2;
        }

        body {
            background: linear-gradient(135deg, var(--primary-color) 0%, var(--secondary-color) 100%);
            min-height: 100vh;
            font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif;
        }

        .chat-container {
            max-width: 900px;
            margin: 0 auto;
            height: 100vh;
            display: flex;
            flex-direction: column;
        }

        .header {
            background: rgba(255, 255, 255, 0.95);
            backdrop-filter: blur(10px);
            padding: 20px;
            border-radius: 15px;
            margin-bottom: 20px;
            box-shadow: 0 5px 15px rgba(0,0,0,0.1);
        }

        .messages {
            flex: 1;
            overflow-y: auto;
            padding: 20px;
            background: rgba(255, 255, 255, 0.9);
            backdrop-filter: blur(10px);
            border-radius: 15px;
            margin-bottom: 20px;
            box-shadow: 0 5px 15px rgba(0,0,0,0.1);
        }

        .message {
            margin-bottom: 15px;
            padding: 15px;
            border-radius: 15px;
            max-width: 80%;
            animation: fadeIn 0.3s ease-in;
        }

        @keyframes fadeIn {
            from { opacity: 0; transform: translateY(10px); }
            to { opacity: 1; transform: translateY(0); }
        }

        .message.user {
            background: var(--primary-color);
            color: white;
            margin-left: auto;
            text-align: right;
        }

        .message.assistant {
            background: #f8f9fa;
            color: #333;
            border-left: 4px solid var(--primary-color);
        }

        .input-container {
            background: rgba(255, 255, 255, 0.95);
            backdrop-filter: blur(10px);
            padding: 20px;
            border-radius: 15px;
            box-shadow: 0 5px 15px rgba(0,0,0,0.1);
        }

        .input-group {
            display: flex;
            gap: 10px;
        }

        .form-control {
            border-radius: 25px;
            border: 2px solid #e9ecef;
            padding: 12px 20px;
        }

        .form-control:focus {
            border-color: var(--primary-color);
            box-shadow: 0 0 0 0.2rem rgba(102, 126, 234, 0.25);
        }

        .btn-primary {
            border-radius: 25px;
            padding: 12px 30px;
            background: var(--primary-color);
            border: none;
            transition: all 0.3s;
        }

        .btn-primary:hover {
            background: var(--secondary-color);
            transform: translateY(-2px);
        }

        .typing-indicator {
            display: inline-flex;
            align-items: center;
            gap: 4px;
        }

        .typing-indicator span {
            height: 8px;
            width: 8px;
            background-color: #667eea;
            border-radius: 50%;
            display: inline-block;
            animation: bounce 1.4s infinite ease-in-out both;
        }

        .typing-indicator span:nth-child(1) {
            animation-delay: -0.32s;
        }

        .typing-indicator span:nth-child(2) {
            animation-delay: -0.16s;
        }

        @keyframes bounce {
            0%, 80%, 100% {
                transform: scale(0);
            }
            40% {
                transform: scale(1);
            }
        }

        .status-badge {
            position: fixed;
            top: 20px;
            right: 20px;
            padding: 8px 16px;
            background: rgba(255, 255, 255, 0.9);
            border-radius: 20px;
            font-size: 14px;
            box-shadow: 0 2px 10px rgba(0,0,0,0.1);
        }

        .status-badge.online {
            color: #28a745;
        }

        .mark {
            background-color: #ffff00;
            padding: 2px 4px;
            border-radius: 3px;
        }
    </style>
</head>
<body>
    <div class="status-badge online" id="statusBadge">
        <span class="spinner-border spinner-border-sm me-2"></span>
        连接中...
    </div>

    <div class="container-fluid p-4">
        <div class="chat-container">
            <div class="header">
                <div class="d-flex align-items-center">
                    <div class="me-3">
                        <span style="font-size: 2rem;">🤖</span>
                    </div>
                    <div>
                        <h1 class="h3 mb-1">美易在线答疑</h1>
                        <p class="text-muted mb-0">美纳里尼中国合规智能助手</p>
                    </div>
                </div>
            </div>

            <div class="messages" id="messages">
                <div class="message assistant">
                    <strong>欢迎使用美易在线答疑！</strong><br><br>
                    我可以帮您解答合规和报销相关问题。请输入您的问题，比如：
                    <br>• 为什么重置密码的邮件一直收不到？
                    <br>• 报销流程是什么？
                    <br>• 如何登录系统？
                    <br><br>
                    <small class="text-muted">系统已加载 {{ qa_count }} 个问答对</small>
                </div>
            </div>

            <div class="input-container">
                <form id="chatForm" class="input-group">
                    <input type="text" class="form-control" id="userInput"
                           placeholder="请输入您的问题..." required>
                    <button type="submit" class="btn btn-primary" id="sendBtn">
                        发送
                    </button>
                </form>
            </div>
        </div>
    </div>

    <script>
        const messagesContainer = document.getElementById('messages');
        const chatForm = document.getElementById('chatForm');
        const userInput = document.getElementById('userInput');
        const sendBtn = document.getElementById('sendBtn');
        const statusBadge = document.getElementById('statusBadge');

        let isTyping = false;

        // 检查连接状态
        async function checkStatus() {
            try {
                const response = await fetch('/api/status');
                const data = await response.json();

                if (data.status === 'ok') {
                    statusBadge.innerHTML = '<span class="fas fa-circle text-success me-2"></span>在线';
                    statusBadge.classList.add('online');
                }
            } catch (error) {
                statusBadge.innerHTML = '<span class="fas fa-circle text-danger me-2"></span>离线';
            }
        }

        // 定期检查状态
        setInterval(checkStatus, 5000);
        checkStatus();

        function addMessage(content, type) {
            const messageDiv = document.createElement('div');
            messageDiv.className = `message ${type}`;
            messageDiv.innerHTML = content;
            messagesContainer.appendChild(messageDiv);
            messagesContainer.scrollTop = messagesContainer.scrollHeight;
        }

        function showTyping() {
            if (isTyping) return;
            isTyping = true;

            const typingDiv = document.createElement('div');
            typingDiv.className = 'message assistant';
            typingDiv.id = 'typingMessage';
            typingDiv.innerHTML = `
                <div class="typing-indicator">
                    <span></span>
                    <span></span>
                    <span></span>
                </div>
                <small class="text-muted ms-2">正在搜索知识库...</small>
            `;
            messagesContainer.appendChild(typingDiv);
            messagesContainer.scrollTop = messagesContainer.scrollHeight;
        }

        function removeTyping() {
            const typingMessage = document.getElementById('typingMessage');
            if (typingMessage) {
                typingMessage.remove();
            }
            isTyping = false;
        }

        async function sendMessage() {
            const message = userInput.value.trim();
            if (!message || isTyping) return;

            // 添加用户消息
            addMessage(message, 'user');
            userInput.value = '';
            sendBtn.disabled = true;

            // 显示输入状态
            showTyping();

            try {
                // 调用API
                const response = await fetch('/api/ask', {
                    method: 'POST',
                    headers: {
                        'Content-Type': 'application/json',
                    },
                    body: JSON.stringify({ question: message })
                });

                const data = await response.json();

                // 移除输入状态
                removeTyping();

                // 添加回答
                if (data.success) {
                    addMessage(data.answer, 'assistant');
                } else {
                    addMessage(data.answer, 'assistant');
                }

            } catch (error) {
                removeTyping();
                addMessage('抱歉，服务暂时不可用。请稍后重试。', 'assistant');
            } finally {
                sendBtn.disabled = false;
                userInput.focus();
            }
        }

        // 处理表单提交
        chatForm.addEventListener('submit', (e) => {
            e.preventDefault();
            sendMessage();
        });

        // 处理回车键
        userInput.addEventListener('keypress', (e) => {
            if (e.key === 'Enter' && !e.shiftKey) {
                e.preventDefault();
                sendMessage();
            }
        });

        // 页面加载完成后的初始化
        document.addEventListener('DOMContentLoaded', () => {
            userInput.focus();
        });
    </script>
    <link rel="stylesheet" href="https://cdnjs.cloudflare.com/ajax/libs/font-awesome/6.0.0/css/all.min.css">
</body>
</html>'''

    with open('templates/index.html', 'w', encoding='utf-8') as f:
        f.write(index_html.replace('{{ qa_count }}', '0'))

def main():
    """主函数"""
    import argparse

    # 创建模板
    create_template()

    # 加载知识库
    load_knowledge_base()

    parser = argparse.ArgumentParser(description='美易机器人Web应用')
    parser.add_argument('--host', default='0.0.0.0', help='主机地址 (默认: 0.0.0.0)')
    parser.add_argument('--port', type=int, default=5000, help='端口号 (默认: 5000)')
    parser.add_argument('--debug', action='store_true', help='调试模式')

    args = parser.parse_args()

    print("Starting web application...")
    print(f"Local access: http://localhost:{args.port}")
    print(f"External access: http://{args.host}:{args.port}")
    print(f"Knowledge base Q&A pairs: {len(qa_pairs)}")
    print("Press Ctrl+C to stop server")

    app.run(host=args.host, port=args.port, debug=args.debug)

if __name__ == "__main__":
    main()